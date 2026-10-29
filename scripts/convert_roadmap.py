"""Convert Java_Roadmap_5_Playlists.xlsx into src/data/roadmap.json.

Usage (from the project root):
    python scripts/convert_roadmap.py [path/to/workbook.xlsx]

Requires openpyxl (pip install openpyxl).
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "Java_Roadmap_5_Playlists.xlsx"
OUT_PATH = ROOT / "src" / "data" / "roadmap.json"

PLAYLIST_COLUMNS = ["G", "H", "I", "J", "K"]
PLAYLIST_KEYS = ["A", "B", "C", "D", "E"]
HEADER_ROW = 3
FIRST_DATA_ROW = 4

EXPECTED_TOPICS = 141
EXPECTED_VIDEOS = 472


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_playlists(ws):
    """Playlist metadata from the Roadmap row-3 headers (G-K) and their hyperlinks.

    Header format: "A · Channel\\nPlaylist name\\nNN videos".
    """
    playlists = []
    for col, expected_key in zip(PLAYLIST_COLUMNS, PLAYLIST_KEYS):
        cell = ws[f"{col}{HEADER_ROW}"]
        lines = [line.strip() for line in clean(cell.value).split("\n") if line.strip()]
        match = re.match(r"^([A-E])\s*·\s*(.+)$", lines[0])
        if not match or match.group(1) != expected_key:
            raise ValueError(f"Unexpected playlist header in {col}{HEADER_ROW}: {cell.value!r}")
        if not cell.hyperlink or not cell.hyperlink.target:
            raise ValueError(f"Playlist header {col}{HEADER_ROW} has no hyperlink")
        header_count = None
        if len(lines) > 2:
            count_match = re.match(r"^(\d+)\s+videos?$", lines[2])
            header_count = int(count_match.group(1)) if count_match else None
        playlists.append(
            {
                "key": expected_key,
                "channel": match.group(2).strip(),
                "name": lines[1] if len(lines) > 1 else "",
                "url": cell.hyperlink.target,
                "videoCount": 0,  # recomputed from the Playlist Index below
                "_headerCount": header_count,
            }
        )
    return playlists


def parse_topics(ws):
    topics = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        topic_id = ws[f"A{row}"].value
        if topic_id is None or clean(ws[f"C{row}"].value) == "":
            continue
        covered_by = [
            key
            for col, key in zip(PLAYLIST_COLUMNS, PLAYLIST_KEYS)
            if clean(ws[f"{col}{row}"].value)
        ]
        topics.append(
            {
                "id": int(topic_id),
                "phase": clean(ws[f"B{row}"].value),
                "topic": clean(ws[f"C{row}"].value),
                "subtopics": clean(ws[f"D{row}"].value),
                "priority": clean(ws[f"E{row}"].value),
                "gapNote": clean(ws[f"L{row}"].value),
                "videos": {key: [] for key in PLAYLIST_KEYS},
                "_coveredBy": covered_by,
            }
        )
    return topics


def parse_videos(ws):
    """Rows of the Playlist Index sheet (header on row 1)."""
    videos = []
    for row in range(2, ws.max_row + 1):
        key = clean(ws[f"A{row}"].value)
        if not key:
            continue
        title_cell = ws[f"D{row}"]
        if not title_cell.hyperlink or not title_cell.hyperlink.target:
            raise ValueError(f"Playlist Index D{row} has no hyperlink")
        topic_ids = [int(part) for part in re.split(r"[,\s]+", clean(ws[f"G{row}"].value)) if part]
        minutes = ws[f"F{row}"].value
        videos.append(
            {
                "key": key,
                "pos": int(ws[f"C{row}"].value),
                "title": clean(title_cell.value),
                "duration": clean(ws[f"E{row}"].value),
                "minutes": float(minutes) if minutes is not None else 0.0,
                "url": title_cell.hyperlink.target,
                "topicIds": topic_ids,
            }
        )
    return videos


def parse_gap_sheet(ws):
    """{topicId: type} from 'Gaps - To Search' (header on row 4)."""
    gaps = {}
    for row in range(5, ws.max_row + 1):
        topic_id = ws[f"A{row}"].value
        if topic_id is None:
            continue
        gaps[int(topic_id)] = clean(ws[f"E{row}"].value)
    return gaps


def main():
    xlsx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = load_workbook(xlsx_path)

    playlists = parse_playlists(wb["Roadmap"])
    topics = parse_topics(wb["Roadmap"])
    videos = parse_videos(wb["Playlist Index"])
    gaps = parse_gap_sheet(wb["Gaps - To Search"])

    topics_by_id = {t["id"]: t for t in topics}
    playlist_by_key = {p["key"]: p for p in playlists}
    errors = []
    warnings = []

    # Attach every video to each topic listed in its "Roadmap #" column.
    for video in videos:
        if video["key"] not in playlist_by_key:
            errors.append(f"Video {video['key']}-{video['pos']} has unknown playlist key")
            continue
        playlist_by_key[video["key"]]["videoCount"] += 1
        if not video["topicIds"]:
            errors.append(f"Video {video['key']}-{video['pos']} is not mapped to any topic")
        for topic_id in video["topicIds"]:
            topic = topics_by_id.get(topic_id)
            if topic is None:
                errors.append(f"Video {video['key']}-{video['pos']} references missing topic #{topic_id}")
                continue
            topic["videos"][video["key"]].append(
                {k: video[k] for k in ("pos", "title", "duration", "minutes", "url")}
            )

    for topic in topics:
        for key in PLAYLIST_KEYS:
            topic["videos"][key].sort(key=lambda v: v["pos"])

    # --- Verification ---
    if len(topics) != EXPECTED_TOPICS:
        errors.append(f"Expected {EXPECTED_TOPICS} topics, found {len(topics)}")
    if len(topics_by_id) != len(topics):
        errors.append("Duplicate topic IDs in Roadmap sheet")
    if len(videos) != EXPECTED_VIDEOS:
        errors.append(f"Expected {EXPECTED_VIDEOS} videos, found {len(videos)}")
    unique_video_ids = {(v["key"], v["pos"]) for v in videos}
    if len(unique_video_ids) != len(videos):
        errors.append("Duplicate (playlist, position) pairs in Playlist Index")

    attached = {
        (key, v["pos"]) for t in topics for key, vids in t["videos"].items() for v in vids
    }
    orphans = unique_video_ids - attached
    if orphans:
        errors.append(f"{len(orphans)} videos not attached to any topic: {sorted(orphans)[:10]}")

    for playlist in playlists:
        header_count = playlist.pop("_headerCount")
        if header_count is not None and header_count != playlist["videoCount"]:
            warnings.append(
                f"Playlist {playlist['key']} header says {header_count} videos, index has {playlist['videoCount']}"
            )

    for topic in topics:
        covered_by = topic.pop("_coveredBy")
        from_index = [key for key in PLAYLIST_KEYS if topic["videos"][key]]
        if covered_by != from_index:
            warnings.append(
                f"Topic #{topic['id']}: Roadmap G-K columns say {covered_by}, Playlist Index says {from_index}"
            )
        gap_type = gaps.get(topic["id"])
        if gap_type == "Not covered" and from_index:
            warnings.append(f"Topic #{topic['id']} is listed as 'Not covered' but has videos")
        if not from_index and not topic["gapNote"]:
            warnings.append(f"Topic #{topic['id']} has no videos and no gap note")

    for message in warnings:
        print(f"WARNING: {message}")
    if errors:
        for message in errors:
            print(f"ERROR: {message}")
        sys.exit(1)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUT_PATH.open("w", encoding="utf-8") as f:
        json.dump({"playlists": playlists, "topics": topics}, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print(
        f"OK: wrote {OUT_PATH.relative_to(ROOT)} with {len(playlists)} playlists, "
        f"{len(topics)} topics, {len(videos)} videos (all attached to at least one topic)."
    )


if __name__ == "__main__":
    main()
